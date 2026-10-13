from __future__ import annotations

from app.ai.teacher_judge.schemas import TeacherJudgeRubricItem
from app.ai.teacher_judge.service import normalize_items_for_export


def test_teacher_judge_normalizes_ai_returned_items() -> None:
    items = normalize_items_for_export(
        [
            {
                "name": "Port 80",
                "desc": "檢查 Web 服務",
                "is_checked": "yes",
                "detectable": "AUTO",
                "detection": "TCP Port 80 探測",
                "suggestion": "請學生確認防火牆設定",
            },
            {"title": "程式碼品質", "detectable": "unknown"},
        ]
    )

    assert items == [
        TeacherJudgeRubricItem(
            id="item-1",
            title="Port 80",
            checked=True,
            detectable="auto",
            detection_method="TCP Port 80 探測",
            fallback="請學生確認防火牆設定",
        ),
        TeacherJudgeRubricItem(
            id="item-2",
            title="程式碼品質",
            checked=False,
            detectable="manual",
            detection_method=None,
            fallback=None,
        ),
    ]


def test_teacher_judge_export_strips_control_characters() -> None:
    """終端輸出的 ANSI 色碼或 PDF 換頁符不可讓整份 Excel 匯出失敗。"""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.ai.teacher_judge.export import export_to_excel

    content = export_to_excel(
        [
            TeacherJudgeRubricItem(
                id="item-1",
                title="log\x1b[31m red",
                detection_method="=cmd|' /C calc'!A0\x0c",
            )
        ],
        summary="page\x0cbreak",
    )

    sheet = load_workbook(BytesIO(content)).active
    assert sheet.cell(row=2, column=2).value == "log[31m red"
    assert sheet.cell(row=2, column=6).value == "'=cmd|' /C calc'!A0"
