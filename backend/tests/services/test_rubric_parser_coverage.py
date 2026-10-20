"""Coverage tests for app.services.rubric_parser (text / docx / pdf / doc parsing)."""

from __future__ import annotations

import io
import struct
import subprocess
import sys
import types
import zipfile
import zlib
from pathlib import Path

import pytest
from docx import Document

from app.services import rubric_parser
from app.services.rubric_parser import (
    _clean_cell,
    _guard_zip_document,
    _markdown_table,
    _terminate_process_tree,
    parse_doc,
    parse_document,
    parse_pdf,
    parse_text,
)


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Rubric title")
    doc.add_paragraph("   ")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Line\nbreak"
    table.cell(1, 1).text = "10"
    doc.add_paragraph("Footer")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ── zip guard ──────────────────────────────────────────────


def test_guard_zip_rejects_corrupted_document() -> None:
    with pytest.raises(ValueError):
        _guard_zip_document(b"definitely not a zip")


def test_guard_zip_rejects_total_size_over_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("a.txt", b"x" * 100)
    monkeypatch.setattr(rubric_parser, "MAX_ZIP_UNCOMPRESSED_BYTES", 10)
    with pytest.raises(ValueError):
        _guard_zip_document(buffer.getvalue())


def _deflate_bomb(real_size: int, declared_size: int | None = None) -> bytes:
    """A single-member DEFLATE archive of ``real_size`` zero bytes.

    With ``declared_size`` the central directory and local header are patched
    to claim that smaller size, with a CRC matching the truncated content, so a
    guard that trusts ``ZipInfo.file_size`` (or ``zipfile``'s own truncating
    reader) sees a tiny, valid-looking member.
    """

    payload = b"\0" * real_size
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", payload)
    data = bytearray(buffer.getvalue())
    if declared_size is None:
        return bytes(data)
    crc = zlib.crc32(payload[:declared_size]) & 0xFFFFFFFF
    local = data.find(b"PK\x03\x04")
    central = data.find(b"PK\x01\x02")
    assert local >= 0 and central >= 0
    struct.pack_into("<I", data, local + 14, crc)
    struct.pack_into("<I", data, local + 22, declared_size)
    struct.pack_into("<I", data, central + 16, crc)
    struct.pack_into("<I", data, central + 24, declared_size)
    return bytes(data)


def test_guard_zip_rejects_honest_deflate_bomb_by_ratio() -> None:
    # 200 KB of zeros deflates to a few hundred bytes: under the size cap but
    # far above the 100x ratio cap
    with pytest.raises(ValueError):
        _guard_zip_document(_deflate_bomb(200 * 1024))


def test_guard_zip_does_not_trust_declared_member_sizes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(rubric_parser, "MAX_ZIP_UNCOMPRESSED_BYTES", 1024 * 1024)
    bomb = _deflate_bomb(4 * 1024 * 1024, declared_size=1024)
    with zipfile.ZipFile(io.BytesIO(bomb)) as archive:
        declared = sum(info.file_size for info in archive.infolist())
        # zipfile itself truncates to the declared size and sees a valid CRC
        assert declared == 1024
        assert len(archive.read("word/document.xml")) == 1024
    with pytest.raises(ValueError, match="document_too_large|MB"):
        _guard_zip_document(bomb)


def test_guard_zip_accepts_small_deflated_member() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("word/document.xml", b"<w:document>hello</w:document>" * 20)
    _guard_zip_document(buffer.getvalue())


def test_guard_zip_rejects_unsupported_compression_method() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_BZIP2) as archive:
        archive.writestr("word/document.xml", b"<w:document/>")
    with pytest.raises(ValueError):
        _guard_zip_document(buffer.getvalue())


# ── docx ───────────────────────────────────────────────────


def test_parse_document_dispatches_docx() -> None:
    assert "Rubric title" in parse_document("Rubric.DOCX", _docx_bytes())


def test_clean_cell_and_markdown_table_helpers() -> None:
    assert _clean_cell(None) == ""
    assert _clean_cell("  a\nb ") == "a b"
    assert _markdown_table([]) == []
    assert _markdown_table([["h"]]) == ["| h |", "| --- |"]


# ── text ───────────────────────────────────────────────────


def test_parse_text_strips_bom_and_rejects_non_utf8() -> None:
    assert parse_text(b"\xef\xbb\xbfhello") == "hello"
    with pytest.raises(ValueError):
        parse_text(b"\xff\xfe\xfa")


def test_parse_document_text_and_unsupported() -> None:
    assert parse_document("notes.md", b"# hi") == "# hi"
    assert parse_document("notes.txt", b"plain") == "plain"
    with pytest.raises(ValueError):
        parse_document("sheet.xlsx", b"")


# ── pdf ────────────────────────────────────────────────────


class _FakeTable:
    def __init__(self, data, bbox) -> None:
        self._data = data
        self.bbox = bbox

    def extract(self):
        return self._data


class _FakePage:
    def __init__(self, tables, words) -> None:
        self._tables = tables
        self._words = words

    def find_tables(self):
        return self._tables

    def extract_words(self):
        return self._words


class _FakePdf:
    def __init__(self, pages) -> None:
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *exc) -> None:
        return None


def _install_fake_pdfplumber(monkeypatch: pytest.MonkeyPatch, pages) -> None:
    module = types.ModuleType("pdfplumber")
    module.open = lambda _stream: _FakePdf(pages)  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "pdfplumber", module)


def test_parse_pdf_tables_and_text_outside_tables(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    table = _FakeTable([["Item", None], ["A\nB", "5"]], (0.0, 0.0, 100.0, 100.0))
    empty_table = _FakeTable([], (0.0, 0.0, 1.0, 1.0))
    words = [
        # Inside the table bbox: excluded.
        {"text": "inside", "x0": 10, "x1": 20, "top": 10, "bottom": 20},
        # line_number given.
        {"text": "second", "x0": 300, "x1": 320, "top": 150, "line_number": 1},
        {"text": "first", "x0": 200, "x1": 220, "top": 150, "line_number": 1},
        # No line_number: falls back to rounded top.
        {"text": "next", "x0": 200, "x1": 220, "top": 199.6, "bottom": 210},
        # Whitespace-only text on its own line: dropped.
        {"text": "   ", "x0": 200, "x1": 220, "top": 400, "bottom": 410},
    ]
    blank_page = _FakePage([], None)
    _install_fake_pdfplumber(
        monkeypatch, [_FakePage([empty_table, table], words), blank_page]
    )

    text = parse_pdf(b"%PDF")
    assert text.split("\n") == [
        "| Item |  |",
        "| --- | --- |",
        "| A B | 5 |",
        "first second",
        "next",
    ]


def test_parse_document_dispatches_pdf(monkeypatch: pytest.MonkeyPatch) -> None:
    _install_fake_pdfplumber(
        monkeypatch,
        [_FakePage([], [{"text": "hello", "x0": 1, "x1": 2, "top": 3}])],
    )
    assert parse_document("r.pdf", b"%PDF") == "hello"


def test_parse_pdf_rejects_too_many_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(rubric_parser, "MAX_PDF_PAGES", 1)
    _install_fake_pdfplumber(monkeypatch, [_FakePage([], []), _FakePage([], [])])
    with pytest.raises(ValueError):
        parse_pdf(b"%PDF")


def test_parse_pdf_without_pdfplumber_returns_503(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from fastapi import HTTPException

    monkeypatch.setitem(sys.modules, "pdfplumber", None)
    with pytest.raises(HTTPException) as exc_info:
        parse_pdf(b"%PDF")
    assert exc_info.value.status_code == 503


# ── doc (LibreOffice) ──────────────────────────────────────


class _FakeProcess:
    def __init__(
        self,
        args,
        *,
        write_output: bool = True,
        returncode: int = 0,
        timeouts: int = 0,
        kwargs=None,
    ) -> None:
        self.args = args
        self.kwargs = kwargs or {}
        self.pid = 4321
        self.returncode = returncode
        self._write_output = write_output
        self._timeouts = timeouts
        self.killed = False

    def communicate(self, timeout=None):
        if self._timeouts > 0:
            self._timeouts -= 1
            raise subprocess.TimeoutExpired(self.args, timeout)
        if self._write_output:
            outdir = Path(self.args[self.args.index("--outdir") + 1])
            (outdir / "source.txt").write_bytes("轉出的文字".encode())
        return b"", b""

    def kill(self) -> None:
        self.killed = True


def _patch_popen(monkeypatch: pytest.MonkeyPatch, **options) -> list[_FakeProcess]:
    created: list[_FakeProcess] = []

    def _popen(args, **kwargs):
        process = _FakeProcess(args, kwargs=kwargs, **options)
        created.append(process)
        return process

    monkeypatch.setattr(rubric_parser.subprocess, "Popen", _popen)
    monkeypatch.setattr(
        rubric_parser.shutil,
        "which",
        lambda name: "/usr/bin/soffice" if name == "soffice" else None,
    )
    return created


def test_parse_doc_success_with_new_session(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(rubric_parser.os, "setsid", lambda: None, raising=False)
    created = _patch_popen(monkeypatch)

    assert parse_document("legacy.doc", b"doc-bytes") == "轉出的文字"
    assert created[0].kwargs.get("start_new_session") is True
    assert created[0].args[0] == "/usr/bin/soffice"


def test_parse_doc_failure_when_no_output(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_popen(monkeypatch, write_output=False)
    with pytest.raises(ValueError):
        parse_doc(b"doc-bytes")


def test_parse_doc_timeout_kills_process(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delattr(rubric_parser.os, "killpg", raising=False)
    monkeypatch.delattr(rubric_parser.os, "getpgid", raising=False)
    created = _patch_popen(monkeypatch, timeouts=2)
    with pytest.raises(ValueError):
        parse_doc(b"doc-bytes")
    assert created[0].killed is True


def test_terminate_process_tree_uses_killpg(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[tuple[int, int]] = []
    # Windows has no SIGKILL; pin a value so the killpg path runs everywhere.
    monkeypatch.setattr(rubric_parser.signal, "SIGKILL", 9, raising=False)
    monkeypatch.setattr(rubric_parser.os, "getpgid", lambda pid: pid + 1, raising=False)
    monkeypatch.setattr(
        rubric_parser.os,
        "killpg",
        lambda pgid, sig: calls.append((pgid, sig)),
        raising=False,
    )
    process = _FakeProcess(["x"])
    _terminate_process_tree(process)  # type: ignore[arg-type]
    assert calls == [(4322, 9)]
    assert process.killed is False
